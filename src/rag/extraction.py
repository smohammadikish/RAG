from __future__ import annotations

import hashlib
import json
import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import Protocol

from rag.exceptions import DocumentLoadError
from rag.models import Document, Metadata
import csv
from pypdf import PdfReader
from docx import Document as DOCXDocument
from openpyxl import load_workbook

SUPPORTED_EXTRACTION_EXTENSIONS = frozenset(
    {".pdf", ".txt", ".md", ".csv", ".json", ".docx", ".xlsx"}
)
DEFAULT_LINE_GROUP_SIZE = 80
DEFAULT_ROW_GROUP_SIZE = 25
MAX_JSON_SEGMENT_CHARS = 2_000

UNSUPPORTED_EXTRACTION_FORMAT = "UNSUPPORTED_EXTRACTION_FORMAT"
TEXT_DECODE_FAILED = "TEXT_DECODE_FAILED"
MALFORMED_CSV = "MALFORMED_CSV"
MALFORMED_JSON = "MALFORMED_JSON"
MALFORMED_DOCX = "MALFORMED_DOCX"
MALFORMED_XLSX = "MALFORMED_XLSX"
ENCRYPTED_PDF = "ENCRYPTED_PDF"
EMPTY_DOCUMENT = "EMPTY_DOCUMENT"
EXTRACTION_FAILED = "EXTRACTION_FAILED"


@dataclass(frozen=True)
class ExtractionResult:
    """Normalized extracted document segments ready for chunking."""

    documents: tuple[Document, ...]
    source: str
    format: str


class DocumentExtractor(Protocol):
    def extract_path(
        self,
        path: str | Path,
        *,
        source_id: str | None = None,
        source: str | None = None,
    ) -> ExtractionResult:
        """Extract RAG documents from a local parser path."""

    def extract_bytes(
        self,
        content: bytes,
        *,
        filename: str,
        source_id: str | None = None,
        source: str | None = None,
    ) -> ExtractionResult:
        """Extract RAG documents from in-memory file bytes."""


class DocumentExtractionError(DocumentLoadError):
    """Safe structured extraction failure."""

    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code
        self.safe_message = message[:512]


class DefaultDocumentExtractor:
    """Extract all public upload formats through one deterministic boundary."""

    def extract_path(
        self,
        path: str | Path,
        *,
        source_id: str | None = None,
        source: str | None = None,
    ) -> ExtractionResult:
        normalized_path = Path(path).expanduser()
        extension = normalized_path.suffix.lower()
        if extension not in SUPPORTED_EXTRACTION_EXTENSIONS:
            raise DocumentExtractionError(
                UNSUPPORTED_EXTRACTION_FORMAT,
                f"Extraction is not available for {extension or 'unknown'} files.",
            )

        source_name = source or str(normalized_path)
        base_metadata = _base_metadata(
            source=source_name,
            filename=normalized_path.name,
            extension=extension,
            source_id=source_id,
        )

        try:
            if extension in {".txt", ".md"}:
                documents = self._extract_text(normalized_path, base_metadata)
            elif extension == ".pdf":
                documents = self._extract_pdf(normalized_path, base_metadata)
            elif extension == ".csv":
                documents = self._extract_csv(normalized_path, base_metadata)
            elif extension == ".json":
                documents = self._extract_json(normalized_path, base_metadata)
            elif extension == ".docx":
                documents = self._extract_docx(normalized_path, base_metadata)
            else:
                documents = self._extract_xlsx(normalized_path, base_metadata)
        except DocumentExtractionError:
            raise
        except FileNotFoundError:
            raise
        except IsADirectoryError:
            raise
        except UnicodeDecodeError as exc:
            raise DocumentExtractionError(TEXT_DECODE_FAILED, "Text must use UTF-8 encoding.") from exc
        except Exception as exc:
            code = {".csv": MALFORMED_CSV, ".json": MALFORMED_JSON,
                    ".docx": MALFORMED_DOCX, ".xlsx": MALFORMED_XLSX}.get(extension, EXTRACTION_FAILED)
            raise DocumentExtractionError(code, f"Malformed {extension.lstrip('.').upper()} file: {source_name}") from exc

        if not documents:
            raise DocumentExtractionError(
                EMPTY_DOCUMENT,
                f"Document contains no extractable text: {source_name}",
            )

        return ExtractionResult(
            documents=tuple(documents),
            source=source_name,
            format=extension.lstrip("."),
        )

    def extract_bytes(
        self,
        content: bytes,
        *,
        filename: str,
        source_id: str | None = None,
        source: str | None = None,
    ) -> ExtractionResult:
        extension = Path(filename).suffix.lower()
        # Save uploaded bytes briefly so the same file readers can handle them.
        with tempfile.TemporaryDirectory(prefix="rag-extract-") as directory:
            path = Path(directory) / f"source{extension}"
            path.write_bytes(content)
            result = self.extract_path(
                path,
                source_id=source_id,
                source=source or filename,
            )
            return ExtractionResult(
                documents=tuple(
                    _with_metadata(
                        document,
                        {
                            "filename": filename,
                            "source": source or filename,
                        },
                    )
                    for document in result.documents
                ),
                source=source or filename,
                format=result.format,
            )

    def _extract_text(
        self,
        path: Path,
        base_metadata: Metadata,
    ) -> list[Document]:
        text = path.read_text(encoding="utf-8")

        lines = text.splitlines()
        documents = []
        for group_index, start in enumerate(
            range(0, len(lines), DEFAULT_LINE_GROUP_SIZE)
        ):
            group_lines = lines[start : start + DEFAULT_LINE_GROUP_SIZE]
            segment_text = "\n".join(group_lines).strip()
            if not segment_text:
                continue
            line_start = start + 1
            line_end = start + len(group_lines)
            metadata = {
                **base_metadata,
                "location_type": "line_range",
                "line_start": line_start,
                "line_end": line_end,
            }
            documents.append(
                _document(
                    text=segment_text,
                    metadata=metadata,
                    location=f"lines:{line_start}-{line_end}",
                    group_index=group_index,
                )
            )
        return documents

    def _extract_pdf(self, path: Path, base_metadata: Metadata) -> list[Document]:
        reader = PdfReader(path)
        if reader.is_encrypted:
            raise DocumentExtractionError(ENCRYPTED_PDF, "Encrypted PDF documents are not supported.")
        documents = []
        for number, page in enumerate(reader.pages, 1):
            text = (page.extract_text() or "").strip()
            if text:
                documents.append(_document(
                    text=f"Page {number}\n{text}",
                    metadata={**base_metadata, "location_type": "page", "page_number": number,
                              "page_count": len(reader.pages)},
                    location=f"page:{number}", group_index=number - 1,
                ))
        return documents

    def _extract_csv(self, path: Path, base_metadata: Metadata) -> list[Document]:
        with path.open(encoding="utf-8-sig", newline="") as stream:
            sample = stream.read(8192)
            stream.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
            except csv.Error:
                dialect = csv.excel
            rows = [tuple(row) for row in csv.reader(stream, dialect, strict=True)]
        return self._table(rows, base_metadata, "rows")

    def _extract_json(self, path: Path, base_metadata: Metadata) -> list[Document]:
        data = json.loads(path.read_text(encoding="utf-8"))
        return [] if data in ({}, []) else _json_documents(data, path="$", base_metadata=base_metadata)

    def _extract_docx(self, path: Path, base_metadata: Metadata) -> list[Document]:
        file = DOCXDocument(path)
        documents = []
        for number, paragraph in enumerate(file.paragraphs, 1):
            if paragraph.text.strip():
                documents.append(_document(
                    text=f"Paragraph {number}\n{paragraph.text.strip()}",
                    metadata={**base_metadata, "location_type": "paragraph", "paragraph_index": number},
                    location=f"paragraph:{number}", group_index=len(documents),
                ))
        for number, table in enumerate(file.tables, 1):
            rows = [tuple(cell.text for cell in row.cells) for row in table.rows]
            documents.extend(self._table(rows, {**base_metadata, "location_type": "table_rows",
                                                "table_index": number}, f"table:{number}:rows"))
        return documents

    def _extract_xlsx(self, path: Path, base_metadata: Metadata) -> list[Document]:
        # Read saved cell values; this reader does not calculate Excel formulas.
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            documents = []
            for sheet in workbook:
                rows = [tuple("" if cell is None else str(cell) for cell in row)
                        for row in sheet.iter_rows(values_only=True)]
                documents.extend(self._table(rows, {**base_metadata, "location_type": "sheet_rows",
                                                    "sheet_name": sheet.title}, f"sheet:{sheet.title}:rows"))
            return documents
        finally:
            workbook.close()

    @staticmethod
    def _table(rows, metadata, location):
        if not _has_non_empty_rows(rows):
            return []
        # Use the first row as column names, with defaults for empty names.
        columns = _columns_from_header(rows[0], max(map(len, rows)))
        return _tabular_documents(
            rows=rows[1:] if len(rows) > 1 else rows, columns=columns,
            base_metadata=metadata, location_prefix=location,
            row_offset=2 if len(rows) > 1 else 1,
        )


class SharedDocumentLoader:
    """DocumentLoader adapter over the shared extractor."""

    def __init__(
        self,
        *,
        extractor: DocumentExtractor | None = None,
        source_id: str | None = None,
        source: str | None = None,
    ) -> None:
        self._extractor = extractor or DefaultDocumentExtractor()
        self._source_id = source_id
        self._source = source

    def load(self, path: str | Path) -> list[Document]:
        return list(
            self._extractor.extract_path(
                path,
                source_id=self._source_id,
                source=self._source,
            ).documents
        )


def _base_metadata(
    *,
    source: str,
    filename: str,
    extension: str,
    source_id: str | None,
) -> Metadata:
    metadata: Metadata = {
        "source": source,
        "filename": filename,
        "file_type": extension.lstrip("."),
    }
    if source_id is not None:
        metadata["document_id"] = source_id
    return metadata


def _tabular_documents(
    *,
    rows: list[tuple[str, ...]],
    columns: list[str],
    base_metadata: Metadata,
    location_prefix: str,
    row_offset: int,
    group_offset: int = 0,
) -> list[Document]:
    documents = []
    non_empty_rows = [
        (row_index, row)
        for row_index, row in enumerate(rows, start=row_offset)
        if any(str(cell).strip() for cell in row)
    ]
    for group_index, start in enumerate(
        range(0, len(non_empty_rows), DEFAULT_ROW_GROUP_SIZE)
    ):
        group = non_empty_rows[start : start + DEFAULT_ROW_GROUP_SIZE]
        if not group:
            continue
        row_start = group[0][0]
        row_end = group[-1][0]
        metadata = {
            **base_metadata,
            "location_type": base_metadata.get("location_type", "row_range"),
            "row_start": row_start,
            "row_end": row_end,
            "columns": columns,
        }
        text = _serialize_rows(group, columns)
        documents.append(
            _document(
                text=text,
                metadata=metadata,
                location=f"{location_prefix}:{row_start}-{row_end}",
                group_index=group_offset + group_index,
            )
        )
    return documents


def _serialize_rows(rows: list[tuple[int, tuple[str, ...]]], columns: list[str]) -> str:
    parts = [f"Columns: {', '.join(columns)}"]
    for row_index, row in rows:
        values = []
        for column_index, column in enumerate(columns):
            value = row[column_index] if column_index < len(row) else ""
            values.append(f"{column}={value}")
        parts.append(f"Row {row_index}: " + " | ".join(values))
    return "\n".join(parts)


def _json_documents(
    value: object,
    *,
    path: str,
    base_metadata: Metadata,
    group_offset: int = 0,
) -> list[Document]:
    if _json_is_scalar(value):
        return [
            _json_document(
                value,
                path=path,
                base_metadata=base_metadata,
                group_index=group_offset,
            )
        ]

    if isinstance(value, dict):
        documents = []
        for key, child in value.items():
            child_path = f"{path}.{_json_path_key(key)}"
            documents.extend(
                _json_value_documents(
                    child,
                    path=child_path,
                    key=key,
                    base_metadata=base_metadata,
                    group_offset=group_offset + len(documents),
                )
            )
        return documents

    if isinstance(value, list):
        documents = []
        for index, child in enumerate(value):
            child_path = f"{path}[{index}]"
            documents.extend(
                _json_value_documents(
                    child,
                    path=child_path,
                    key=None,
                    base_metadata=base_metadata,
                    group_offset=group_offset + len(documents),
                )
            )
        return documents

    return []


def _json_value_documents(
    value: object,
    *,
    path: str,
    key: str | None,
    base_metadata: Metadata,
    group_offset: int,
) -> list[Document]:
    serialized = json.dumps(value, ensure_ascii=False, indent=2, sort_keys=True)
    if len(serialized) <= MAX_JSON_SEGMENT_CHARS or _json_is_scalar(value):
        metadata = {**base_metadata, "location_type": "json_path", "json_path": path}
        if key is not None:
            metadata["json_key"] = key
        return [
            _document(
                text=f"JSON path {path}\n{serialized}",
                metadata=metadata,
                location=f"json:{path}",
                group_index=group_offset,
            )
        ]
    return _json_documents(
        value,
        path=path,
        base_metadata=base_metadata,
        group_offset=group_offset,
    )


def _json_document(
    value: object,
    *,
    path: str,
    base_metadata: Metadata,
    group_index: int,
) -> Document:
    metadata = {**base_metadata, "location_type": "json_path", "json_path": path}
    return _document(
        text=f"JSON path {path}\n{json.dumps(value, ensure_ascii=False)}",
        metadata=metadata,
        location=f"json:{path}",
        group_index=group_index,
    )


def _document(
    *,
    text: str,
    metadata: Metadata,
    location: str,
    group_index: int,
) -> Document:
    source = str(metadata["source"])
    document_id = _stable_document_id(source, location, group_index, text)
    return Document(
        id=document_id,
        text=text,
        source=source,
        metadata={**metadata, "source_location": location},
    )


def _stable_document_id(
    source: str,
    location: str,
    group_index: int,
    text: str,
) -> str:
    digest = hashlib.sha256(
        f"{source}:{location}:{group_index}:{text}".encode()
    ).hexdigest()
    return f"doc_{digest[:24]}"


def _columns_from_header(header: tuple[object, ...], column_count: int) -> list[str]:
    columns = []
    for index in range(max(column_count, len(header))):
        raw = header[index] if index < len(header) else ""
        value = str(raw).strip()
        columns.append(value or f"column_{index + 1}")
    return columns


def _has_non_empty_rows(rows: list[tuple[object, ...]]) -> bool:
    return any(any(str(cell).strip() for cell in row) for row in rows)


def _json_is_scalar(value: object) -> bool:
    return value is None or isinstance(value, str | int | float | bool)


def _json_path_key(key: str) -> str:
    if key.replace("_", "").isalnum() and not key[:1].isdigit():
        return key
    return json.dumps(key, ensure_ascii=False)


def _with_metadata(document: Document, metadata: Metadata) -> Document:
    merged = {**document.metadata, **metadata}
    return document.model_copy(update={"source": merged["source"], "metadata": merged})
